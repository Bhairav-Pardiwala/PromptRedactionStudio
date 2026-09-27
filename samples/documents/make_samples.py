"""Regenerate the synthetic sample documents in this folder.

Every value is made up and comes from a range reserved for examples: example.com email
addresses, US 555-01xx and UK 020 7946 0xxx phone numbers, the standard test card
numbers, the IBANs from the IBAN documentation, the SSN range reserved for advertising,
and private IP addresses. The names are fictional, and several are deliberately ordinary
English words (Sample, Placeholder, Test), which makes them hard for the name model.

The generators need three libraries the app itself does not, so install them somewhere
other than the app's environment:

    python -m venv .venv-samples
    .venv-samples/Scripts/pip install python-docx openpyxl reportlab      (Windows)
    .venv-samples/bin/pip install python-docx openpyxl reportlab          (macOS/Linux)
    .venv-samples/Scripts/python samples/documents/make_samples.py

Dates in the content are fixed and document timestamps pinned, so running this again
rewrites the same documents rather than a new variant. Two things still change per run:
the timestamps on the entries inside each .docx/.xlsx zip, and the "modified" time openpyxl
always writes into an .xlsx. Neither is content anyone reads.
"""

from __future__ import annotations

import argparse
import csv
import datetime as dt
from pathlib import Path

from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.shared import Pt
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.comments import Comment
from openpyxl.styles import Font
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.pdfgen import canvas
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

# One fixed moment, used wherever a document would otherwise record "now".
FIXED = dt.datetime(2026, 9, 15, 9, 30)
HYPERLINK = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink"

NAMES = [
    "Jane Doe", "John Roe", "Priya Placeholder", "Marcus Testwell",
    "Elena Sample", "Tomas Fictitious", "Sarah Dummy", "Robert Test",
]
EMAILS = ["jane.doe", "john.roe", "priya.placeholder", "marcus.testwell", "elena.sample"]
CARDS = {"visa": "4111 1111 1111 1111", "mastercard": "5555 5555 5555 4444", "amex": "3782 822463 10005"}
IBANS = ["GB82 WEST 1234 5698 7654 32", "DE89 3704 0044 0532 0130 00"]
NHS_NUMBER = "943 476 5919"


def name(n: int) -> str:
    return NAMES[n % len(NAMES)]


def email(n: int) -> str:
    return EMAILS[n % len(EMAILS)] + "@example.com"


def us_phone(n: int) -> str:
    return "+1 (415) 555-" + format(100 + n % 100, "04d")


def uk_phone(n: int) -> str:
    return "+44 20 7946 09" + str(58 - n % 10)


def address(n: int) -> str:
    return str(42 + n) + " Example Street, Springfield"


def ssn(n: int) -> str:
    return "987-65-" + str(4320 + n % 10)


def pin_docx_dates(doc) -> None:
    doc.core_properties.created = FIXED
    doc.core_properties.modified = FIXED


def pin_xlsx_dates(wb) -> None:
    wb.properties.created = FIXED
    wb.properties.modified = FIXED


# --- Word -------------------------------------------------------------------------------


def complaint_docx(path: Path) -> None:
    """Header, footer, a name split across bold and plain runs, a table, a hyperlink whose
    address hides an email, and author metadata."""
    doc = Document()
    section = doc.sections[0]
    section.header.paragraphs[0].text = "Customer: " + name(1)
    section.footer.paragraphs[0].text = "Contact: " + email(1)

    doc.add_paragraph("Customer Service Complaint", style="Heading 1")
    doc.add_paragraph(
        "Dear Support Team,\n\nI am writing to lodge a formal complaint regarding my recent "
        "transaction. My name is " + name(2) + ", and I can be reached at " + us_phone(1)
        + " or " + email(2) + "."
    )
    doc.add_paragraph(
        "The charge of $950.00 was made on my Visa card (full card number: " + CARDS["visa"]
        + "). This occurred without authorization."
    )

    split = doc.add_paragraph("Please contact ")
    split.add_run("Jane").bold = True
    split.add_run(" Doe in our fraud department for resolution.")

    doc.add_paragraph("Customer Details:", style="Heading 2")
    table = doc.add_table(rows=3, cols=2)
    table.style = "Light Grid Accent 1"
    for row, (field, value) in zip(table.rows, [("Field", "Value"), ("Name", name(3)), ("Email", email(3))]):
        row.cells[0].text = field
        row.cells[1].text = value

    # python-docx has no hyperlink API, so build the element and its relationship by hand.
    link = doc.add_paragraph("For assistance, ")
    rel_id = doc.part.relate_to("mailto:" + email(4), HYPERLINK, is_external=True)
    link._p.append(parse_xml(
        '<w:hyperlink ' + nsdecls("w", "r") + ' r:id="' + rel_id + '"><w:r><w:rPr>'
        '<w:rStyle w:val="Hyperlink"/></w:rPr><w:t>email her</w:t></w:r></w:hyperlink>'
    ))
    link.add_run(".")

    doc.core_properties.author = "Jane Placeholder"
    doc.core_properties.last_modified_by = "John Sample"
    pin_docx_dates(doc)
    doc.save(path)


def hr_review_docx(path: Path) -> None:
    """Names, an employee ID and salary in the body, a team table, and a Word comment
    whose text and author both name people."""
    doc = Document()
    doc.add_paragraph("Performance Review", style="Heading 1")
    body = doc.add_paragraph(
        "Employee: " + name(5) + "\nManager: " + name(6) + "\nEmployee ID: EMP-12350\nSalary: $85,000"
    )

    doc.add_paragraph("Team Members", style="Heading 2")
    table = doc.add_table(rows=6, cols=3)
    table.style = "Light Grid Accent 1"
    for cell, heading in zip(table.rows[0].cells, ["Name", "Email", "Phone"]):
        cell.text = heading
    for i in range(1, 6):
        cells = table.rows[i].cells
        cells[0].text = name(10 + i)
        cells[1].text = email(10 + i)
        cells[2].text = us_phone(i)

    comment = doc.add_comment(
        body.runs, text="Discuss the raise with " + name(4) + " before Friday.",
        author="Sarah Dummy", initials="SD",
    )
    comment._comment_elm.date = FIXED.replace(tzinfo=dt.timezone.utc)  # python-docx stamps "now"
    pin_docx_dates(doc)
    doc.save(path)


def referral_docx(path: Path) -> None:
    """A patient letter: name, date of birth, address, an NHS number, a UK phone."""
    doc = Document()
    doc.add_paragraph("Medical Referral Letter", style="Heading 1")
    doc.add_paragraph(
        "Date: 15 September 2026\n\nPatient Name: " + name(7) + "\nDate of Birth: 15/03/1985\n"
        "Address: " + address(7) + "\nNHS Number: " + NHS_NUMBER + "\n"
    )
    doc.add_paragraph("Referring Doctor: Dr. " + name(8) + "\nContact: " + uk_phone(1))
    doc.add_paragraph(
        "The patient presents with symptoms requiring specialist consultation. "
        "Please see attached notes."
    )
    pin_docx_dates(doc)
    doc.save(path)


def control_docx(path: Path) -> None:
    """No personal data at all: redaction should leave it unchanged."""
    doc = Document()
    doc.add_paragraph("Product Description: Enterprise Software Suite", style="Heading 1")
    doc.add_paragraph(
        "Our enterprise software suite provides comprehensive solutions for business process "
        "automation, data analytics, and customer relationship management."
    )
    doc.add_paragraph("Key Features", style="Heading 2")
    for feature in ["Real-time data processing and analytics", "Cloud-based deployment options",
                    "Multi-tenant architecture", "API-first integration model"]:
        doc.add_paragraph(feature)
    doc.add_paragraph("Technical Specifications", style="Heading 2")
    table = doc.add_table(rows=4, cols=2)
    table.style = "Light Grid Accent 1"
    rows = [("Component", "Details"), ("Runtime", "Python 3.10+"),
            ("Database", "PostgreSQL 14+"), ("Frontend", "React 18+")]
    for row, (left, right) in zip(table.rows, rows):
        row.cells[0].text = left
        row.cells[1].text = right
    pin_docx_dates(doc)
    doc.save(path)


# --- Excel ------------------------------------------------------------------------------


def customer_list_xlsx(path: Path) -> None:
    """Two sheets sharing names, phone numbers stored as numbers, a card number in a text
    cell, a SUM formula, a cell comment naming someone, and a chart."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Customers"
    ws.append(["Name", "Email", "Phone", "City"])
    cities = ["Springfield", "Shelbyville", "Capital City", "Ogdenville"]
    for i in range(1, 11):
        # Rows 3 and 6 hold the phone as a number, the way a spreadsheet import often does.
        phone = 4155550100 + i if i in (2, 5) else us_phone(i)
        ws.append([name(20 + i), email(20 + i), phone, cities[i % 4]])
    ws["D3"] = CARDS["visa"]
    ws["A2"].comment = Comment("Priority customer: call " + name(1) + " back first.", "Sarah Dummy")

    orders = wb.create_sheet("Orders")
    orders.append(["Order ID", "Customer Name", "Amount"])
    for i in range(1, 6):
        orders.append(["ORD-" + str(10000 + i), name(20 + i), 100 + i * 50])
    orders["D1"] = "Total"
    orders["D2"] = "=SUM(C2:C6)"

    for sheet in (ws, orders):
        for cell in sheet[1]:
            cell.font = Font(bold=True)

    chart = BarChart()
    chart.title = "Order Amounts"
    chart.add_data(Reference(orders, min_col=3, min_row=1, max_row=6), titles_from_data=True)
    orders.add_chart(chart, "F2")

    wb.properties.creator = "Jane Placeholder"
    pin_xlsx_dates(wb)
    wb.save(path)


def payroll_xlsx(path: Path) -> None:
    """Names, IBANs, SSNs (reserved range) and salaries in one sheet."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Payroll"
    ws.append(["Employee Name", "IBAN", "SSN", "Salary"])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for i in range(1, 8):
        ws.append([name(30 + i), IBANS[i % 2], ssn(i), 60000 + i * 5000])
    pin_xlsx_dates(wb)
    wb.save(path)


# --- PDF --------------------------------------------------------------------------------


def invoice_pdf(path: Path) -> None:
    """Two pages: billing details on the first, card and IBAN on the second."""
    doc = SimpleDocTemplate(str(path), pagesize=letter, invariant=1)
    styles = getSampleStyleSheet()
    title = ParagraphStyle("InvoiceTitle", parent=styles["Heading1"], fontSize=24,
                           textColor=colors.HexColor("#003366"), spaceAfter=30)
    plain = TableStyle([("FONTNAME", (0, 0), (-1, -1), "Helvetica"), ("FONTSIZE", (0, 0), (-1, -1), 10)])

    billing = Table([["Bill To:"], ["Name: " + name(40)], ["Address: " + address(40)],
                     ["Email: " + email(40)], ["Phone: " + us_phone(40)]], colWidths=[5.5 * inch])
    billing.setStyle(plain)

    items = Table([
        ["Item", "Description", "Qty", "Unit Price", "Total"],
        ["SVC-001", "Consulting Service", "5", "$200.00", "$1,000.00"],
        ["SVC-002", "Support Hours", "10", "$75.00", "$750.00"],
        ["SVC-003", "Technical Review", "1", "$500.00", "$500.00"],
    ], colWidths=[1 * inch, 2 * inch, 0.7 * inch, 1.2 * inch, 1.2 * inch])
    items.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.grey),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.whitesmoke),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), 1, colors.black),
    ]))

    totals = Table([["Subtotal:", "$2,250.00"], ["Tax (10%):", "$225.00"], ["Total:", "$2,475.00"]],
                   colWidths=[4 * inch, 2 * inch])
    totals.setStyle(TableStyle([("ALIGN", (0, 0), (-1, -1), "RIGHT")]))

    payment = Table([["Card Type:", "Visa"], ["Card Number:", CARDS["visa"]], ["IBAN:", IBANS[1]]],
                    colWidths=[1.5 * inch, 3.5 * inch])
    payment.setStyle(plain)

    doc.build([
        Paragraph("INVOICE", title), billing, Spacer(1, 0.3 * inch), items,
        Spacer(1, 0.4 * inch), totals, PageBreak(),
        Paragraph("Payment Information", styles["Heading2"]), Spacer(1, 0.2 * inch), payment,
    ])


def scanned_pdf(path: Path) -> None:
    """Shapes only -- no text layer at all, which is what a scan without OCR looks like."""
    c = canvas.Canvas(str(path), pagesize=letter, invariant=1)
    c.setFillColor(colors.lightgrey)
    c.rect(0.75 * inch, 0.75 * inch, 7 * inch, 9.5 * inch, stroke=0, fill=1)
    c.setFillColor(colors.black)
    y = 9.6 * inch
    for width in (4.5, 6.0, 5.2, 6.1, 3.8, 5.9, 6.0, 4.4, 5.5, 2.9):
        c.rect(1.1 * inch, y, width * inch, 0.12 * inch, stroke=0, fill=1)   # "lines of text"
        y -= 0.35 * inch
    c.rect(5.2 * inch, 1.3 * inch, 1.8 * inch, 0.5 * inch, stroke=1, fill=0)  # "signature box"
    c.showPage()
    c.save()


# --- Text -------------------------------------------------------------------------------


def tickets_csv(path: Path) -> None:
    """UTF-8 with a byte order mark, as Excel writes it. The name appears both in its own
    column and inside the free-text message."""
    with open(path, "w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.writer(handle)
        writer.writerow(["ticket_id", "customer_name", "email", "phone", "message"])
        for i in range(1, 13):
            who = name(50 + i)
            writer.writerow([
                "TKT-" + str(1000 + i), who, email(50 + i), us_phone(i),
                "Customer " + who + " reports an issue. Contact them at " + us_phone(i) + ".",
            ])


def postmortem_md(path: Path) -> None:
    """Markdown with a table, private IP addresses, engineer names, an email and a URL."""
    path.write_text(
        "# Incident Postmortem Report\n\n"
        "**Date:** 2026-09-15\n**Incident ID:** INC-20260915-001\n\n"
        "## Summary\n\nService outage lasting 47 minutes due to a database connectivity issue.\n\n"
        "## Timeline\n\n| Time | Event |\n|------|-------|\n"
        "| 14:30 | Alert triggered by " + name(60) + " |\n"
        "| 14:35 | " + name(61) + " confirmed the issue |\n"
        "| 14:50 | Database reconnected by " + name(62) + " |\n"
        "| 15:17 | All systems operational |\n\n"
        "## Affected Systems\n\n- API Server: 192.168.1.100\n- Database: 192.168.1.101\n"
        "- Cache: 192.168.1.102\n\n"
        "## Engineer Notes\n\n"
        "- " + name(63) + " (engineer@example.com) led the incident response\n"
        "- Contact: " + uk_phone(2) + "\n"
        "- Further details: https://example.com/incident/details\n\n"
        "## Action Items\n\n1. Review database connection pool settings\n"
        "2. Implement automated failover\n3. Update monitoring thresholds\n",
        encoding="utf-8", newline="\n",
    )


def chat_txt(path: Path) -> None:
    """A support chat: name, phone, email and street address in conversation."""
    path.write_text(
        "SUPPORT CHAT TRANSCRIPT\nDate: 2026-09-15 14:22\n\n"
        "[14:22] Customer: Hi, my name is " + name(70) + ". I need help with my account.\n"
        "[14:23] Agent: Hello! How can I help you today?\n"
        "[14:24] Customer: I need to verify my phone number. It's " + us_phone(1) + ".\n"
        "[14:25] Agent: I'll help you with that. Can you also provide your email?\n"
        "[14:26] Customer: Sure, it's " + email(1) + ".\n"
        "[14:27] Agent: And your address for verification?\n"
        "[14:28] Customer: " + address(70) + "\n"
        "[14:29] Agent: Thank you. I've verified your information. Anything else?\n"
        "[14:30] Customer: No, that's all. Thanks!\n",
        encoding="utf-8", newline="\n",
    )


def legacy_cp1252_txt(path: Path) -> None:
    """Windows-1252, not UTF-8: an en dash, curly quotes and an accented name."""
    path.write_text(
        "Project Notes – 2026-09-15\n\n"
        "Contact: José Placeholder\nPhone: " + us_phone(70) + "\nEmail: " + email(70) + "\n\n"
        "“Quote from José about the meeting.”\n\n"
        "Next steps: Contact José at " + address(71) + "\n",
        encoding="cp1252", newline="\r\n",
    )


def corrupt_docx(path: Path) -> None:
    """Not a zip at all, so not a Word file, whatever the extension says."""
    path.write_bytes(bytes(range(8)) * 100)


def unsupported_rtf(path: Path) -> None:
    r"""RTF is not a supported type; the upload should be refused before parsing."""
    path.write_text(
        "{\\rtf1\\ansi\\ansicpg1252\n{\\fonttbl\\f0\\fswiss\\fcharset0 Arial;}\n"
        "\\pard\\plain\\fs20 Unsupported RTF document for " + name(1) + ".\\par\n}\n",
        encoding="ascii", newline="\n",
    )


SAMPLES = [
    ("01-customer-complaint.docx", complaint_docx),
    ("02-hr-performance-review.docx", hr_review_docx),
    ("03-medical-referral.docx", referral_docx),
    ("04-customer-list.xlsx", customer_list_xlsx),
    ("05-payroll.xlsx", payroll_xlsx),
    ("06-invoice.pdf", invoice_pdf),
    ("07-scanned-image-only.pdf", scanned_pdf),
    ("08-support-tickets.csv", tickets_csv),
    ("09-incident-postmortem.md", postmortem_md),
    ("10-chat-transcript.txt", chat_txt),
    ("11-legacy-notes-cp1252.txt", legacy_cp1252_txt),
    ("12-no-pii-control.docx", control_docx),
    ("13-corrupt.docx", corrupt_docx),
    ("14-unsupported.rtf", unsupported_rtf),
]


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    parser.add_argument("--out", type=Path, default=Path(__file__).resolve().parent,
                        help="folder to write into (default: this script's folder)")
    out = parser.parse_args().out
    out.mkdir(parents=True, exist_ok=True)
    for filename, build in SAMPLES:
        build(out / filename)
        print(filename.ljust(32), (out / filename).stat().st_size, "bytes")


if __name__ == "__main__":
    main()
